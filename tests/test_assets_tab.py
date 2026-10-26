"""The Assets page: bringing in a map's hardware (by hand and after each crawl), editing, exporting, and the server
and phone readers."""
import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest  # noqa: E402
from PyQt5.QtCore import QObject, pyqtSignal  # noqa: E402
from PyQt5.QtWidgets import QApplication, QWidget  # noqa: E402
from test_assets import chassis_part, map_with, switch_device  # noqa: E402

from nomad import hostinfo  # noqa: E402
from nomad.ipam.assets import RETIRED, AssetStore  # noqa: E402
from nomad.ipam.store import IpamStore  # noqa: E402
from nomad.netmap import hardware  # noqa: E402
from nomad.netmap.model import PHONE, Host  # noqa: E402
from nomad.ui import assets_tab as assets_module  # noqa: E402
from nomad.ui.asset_dialogs import AssetDialog, AssetImportDialog, ReadDetailsDialog  # noqa: E402
from nomad.ui.assets_tab import COL_PROBLEMS, COL_SERIAL, COL_SLOT, COL_STATUS, AssetsTab  # noqa: E402
from nomad.ui.ipam_history_dialog import HistoryDialog  # noqa: E402
from nomad.ui.ipam_tab import IpamTab  # noqa: E402


@pytest.fixture(scope="module")
def app():
    return QApplication.instance() or QApplication([])


class Navigator:
    def __init__(self):
        self.current = None

    def setCurrentWidget(self, widget):
        self.current = widget


class FakeMapPage(QObject):
    """The Network Map page, as the Assets page uses it."""
    hardware_found = pyqtSignal()
    map_shown = pyqtSignal()

    def __init__(self, network_map):
        super().__init__()
        self.network_map = network_map
        self.worker = None
        self.read_asked = 0
        self.shown = None

    def map_name(self):
        return "Lab map"

    def reading_hardware(self):
        return False

    def read_hardware_again(self, keys=None):
        self.read_asked += 1
        return True

    def show_device(self, key):
        self.shown = key


class Window(QWidget):
    def __init__(self, network_map=None):
        super().__init__()
        self.navigator = Navigator()
        self.netmap_tab = FakeMapPage(network_map)
        self.session_store = None

    def __getattr__(self, name):
        return lambda *args, **kwargs: None


def lab_map():
    psu = dict(chassis_part("LIT1", "PWR-C1-715WAC", "Switch 1 - Power Supply A", index=1010),
               **{"class": hardware.POWER, "chassis": 1000})
    phone = Host(mac="00-AA-BB-CC-DD-EE", device="sw1", port="Gi1/0/5", ip="10.10.0.22", name="SEP00AABBCCDDEE",
                 platform="Cisco IP Phone 8845", software="sip8845.14-1")
    return map_with(switch_device(serials=("FOC1", "FOC2"), parts=[psu]), hosts=[phone])


@pytest.fixture
def page(app, tmp_path):
    window = Window(lab_map())
    store = IpamStore(str(tmp_path / "ipam.db"), user="tester")
    window.ipam_tab = IpamTab(window)
    window.ipam_tab.local_store = store  # So it doesn't open the real database or the tribe's
    tab = AssetsTab(window)
    tab.fill_sources()
    yield tab, store, window
    window.ipam_tab.sync_timer.stop()
    store.close()


def rows(table):
    return [[table.item(row, column).text() for column in range(table.columnCount())]
            for row in range(table.rowCount()) if not table.isRowHidden(row)]


def bring_in(monkeypatch, tab):
    shown = []
    monkeypatch.setattr(AssetImportDialog, "exec_", lambda dialog: shown.append(dialog) or True)
    tab.bring_in_from_map()
    return shown


def test_bring_in_the_maps_hardware(page, monkeypatch):
    tab, store, window = page
    assert "No assets yet" in tab.summary_label.text()
    [dialog] = bring_in(monkeypatch, tab)
    assert dialog.table.rowCount() == 4  # Two stack members, a power supply and a phone
    listed = {row[COL_SERIAL]: row for row in rows(tab.table)}
    assert set(listed) == {"FOC1", "FOC2", "LIT1", ""}
    assert listed["FOC2"][COL_SLOT] == "Switch 2" and listed["LIT1"][0] == "sw1"  # A part shows its chassis' name
    assert "No serial number" in listed[""][COL_PROBLEMS]  # The phone, known by its MAC
    assert "3 devices" in tab.summary_label.text() and "1 parts" in tab.summary_label.text()
    assert "Brought in from the map: 4 new" in tab.status_label.text()
    assert tab.source_combo.currentText().startswith("Local  (4)")


def test_parts_can_be_left_out(page, monkeypatch):
    tab, store, window = page
    tab.parts_check.setChecked(False)
    bring_in(monkeypatch, tab)
    assert sorted(row[COL_SERIAL] for row in rows(tab.table)) == ["", "FOC1", "FOC2"]


def test_each_crawl_is_brought_in_by_itself(page, monkeypatch):
    tab, store, window = page
    window.netmap_tab.hardware_found.emit()
    assets = AssetStore(store)
    assert assets.count() == 4
    # Swapped hardware, and a location typed by hand the map disagrees with
    assets.set_asset(assets.by_serial("FOC1").id, location="Closet 9")
    later = lab_map()
    later.devices["sw1"].parts[1]["serial"] = "FOC9"
    site = later.new_group("HQ")
    later.set_group(["sw1"], site.key)
    window.netmap_tab.network_map = later
    window.netmap_tab.hardware_found.emit()
    assert assets.by_serial("FOC2").status == "replaced" and assets.by_serial("FOC9") is not None
    assert assets.by_serial("FOC1").location == "Closet 9"  # Left for review
    assert tab.review_button.isVisibleTo(tab) and "(1)" in tab.review_button.text()
    shown = []
    monkeypatch.setattr(AssetImportDialog, "exec_", lambda dialog: shown.append(dialog) or True)
    tab.show_review()
    assert shown and not tab.review_button.isVisibleTo(tab)
    tab.auto_check.setChecked(False)
    assets.delete_asset(assets.by_serial("FOC9").id)
    window.netmap_tab.hardware_found.emit()
    assert assets.by_serial("FOC9") is None  # Not brought in by itself when that's off


def test_add_edit_and_set_status(page, monkeypatch):
    tab, store, window = page
    source = tab.source()
    dialog = AssetDialog(tab, source)
    dialog.inputs["hostname"].setText("spare-sw")
    dialog.inputs["serial"].setText("foc77")
    dialog.inputs["model"].setText("C9300-48P")
    dialog.inputs["warranty_end"].setText("not a date")
    dialog.save()
    assert "Write dates like" in dialog.error_label.text()
    dialog.inputs["warranty_end"].setText("2020-01-01")
    dialog.save()
    asset = dialog.result_item
    assert asset.serial == "FOC77" and asset.hostname == "spare-sw"
    tab.fill_table()
    tab.select_ids({asset.id})
    assert "Warranty ended 2020-01-01" in rows(tab.table)[0][COL_PROBLEMS]
    tab.set_many({"status": RETIRED}, "Setting the status")
    assert rows(tab.table)[0][COL_STATUS] == "Retired"
    assert "Warranty ended" not in rows(tab.table)[0][COL_PROBLEMS]  # Out of service: no longer a problem
    edit = AssetDialog(tab, source, AssetStore(store).asset(asset.id))
    edit.inputs["owner"].setText("Network team")
    edit.save()
    assert AssetStore(store).asset(asset.id).owner == "Network team"
    history = HistoryDialog(tab, store, "", "asset", AssetStore(store).asset(asset.id))
    assert history.table.rowCount() == 3
    everything = HistoryDialog(tab, store, "", "assets", None)
    assert everything.table.rowCount() == 3 and "spare-sw" in everything.table.item(0, 2).text()


def test_show_a_map_device_and_back(page, monkeypatch):
    tab, store, window = page
    bring_in(monkeypatch, tab)
    tab.show_map_device("sw1")
    selected = {asset.serial for asset in tab.selected_assets()}
    assert selected == {"FOC1", "FOC2", "LIT1"}
    tab.select_ids({AssetStore(store).by_serial("FOC1").id})
    assert "Show on the Network Map" in tab.details.toHtml() and "Parts (1)" in tab.details.toHtml()
    tab.show_on_map()
    assert window.netmap_tab.shown == "sw1"
    tab.read_hardware_again()
    assert window.netmap_tab.read_asked == 1


def test_export(page, monkeypatch, tmp_path):
    tab, store, window = page
    bring_in(monkeypatch, tab)
    for name in ("assets.csv", "assets.xlsx"):
        path = str(tmp_path / name)
        monkeypatch.setattr(assets_module.QFileDialog, "getSaveFileName", lambda *args, **kwargs: (path, ""))
        tab.export()
        assert os.path.exists(path) and "Exported 4 assets" in tab.status_label.text()
    text = open(tmp_path / "assets.csv", encoding="utf-8-sig").read()
    assert text.splitlines()[0].startswith("Hostname,Slot,Kind,Model") and "FOC2" in text
    from openpyxl import load_workbook
    sheet = load_workbook(tmp_path / "assets.xlsx").active
    assert sheet["A1"].value == "Hostname" and sheet.max_row == 5


def test_server_details_are_reviewed_then_recorded(page, monkeypatch):
    tab, store, window = page
    answers = [hostinfo.HostDetails("10.1.1.10", hostinfo.WINDOWS, hostname="srv1.corp.example", maker="Dell Inc.",
                                    model="PowerEdge R650", serial="7XK2ABC", os="Microsoft Windows Server 2022",
                                    software="10.0.20348")]
    monkeypatch.setattr(ReadDetailsDialog, "exec_", lambda dialog: True)
    monkeypatch.setattr(ReadDetailsDialog, "answers", lambda dialog: answers)
    bring_in(monkeypatch, tab)  # (Makes AssetImportDialog say yes too)
    tab.read_details(["10.1.1.10"])
    server = AssetStore(store).by_serial("7XK2ABC")
    assert server.kind == "server" and server.model == "PowerEdge R650"
    assert server.software == "Microsoft Windows Server 2022 10.0.20348" and server.mgmt_ip == "10.1.1.10"


def test_read_details_dialog_fields(app):
    dialog = ReadDetailsDialog(None, None, ["10.1.1.10"])
    assert dialog.addresses() == ["10.1.1.10"] and dialog.kind() == "server"
    assert not dialog.user_input.isEnabled()  # This Windows login
    dialog.windows_combo.setCurrentIndex(1)
    assert dialog.user_input.isEnabled()
    assert dialog.login() is None and "user name" in dialog.status_label.text()
    dialog.method_combo.setCurrentIndex(dialog.method_combo.findData(hostinfo.PHONE))
    assert dialog.kind() == PHONE and dialog.login() == {}
    dialog.on_answer(hostinfo.HostDetails("10.10.0.22", hostinfo.PHONE, model="CP-8845", serial="FCH1"))
    dialog.on_answer(hostinfo.HostDetails("10.10.0.23", hostinfo.PHONE, error="Its web page didn't answer"))
    assert [details.address for details in dialog.answers()] == ["10.10.0.22"]
    dialog.reject()


# --------------------------------------------------------------------- What servers and phones say of themselves

def test_windows_answer_and_its_errors():
    def runner(arguments, request, timeout):
        assert "-EncodedCommand" in arguments and '"password": "s3cret"' in request  # On stdin, not the command line
        return 0, '{"hostname":"srv1","domain":"corp.example","maker":"Dell Inc.","model":"PowerEdge R650",' \
                  '"serial":"7XK2ABC","os":"Microsoft Windows Server 2022 Standard","version":"10.0.20348",' \
                  '"mac":"00:11:22:33:44:55"}\n', ""
    details = hostinfo.windows_details("10.1.1.10", "CORP\\admin", "s3cret", runner=runner)
    assert details.ok and details.hostname == "srv1.corp.example" and details.serial == "7XK2ABC"
    assert details.mac == "00-11-22-33-44-55"
    denied = hostinfo.windows_details("10.1.1.11", runner=lambda *args: (1, "", "New-CimSession : Access is denied."))
    assert not denied.ok and "administrator" in denied.error
    vm = hostinfo.parse_windows("10.1.1.12", 0, '{"hostname":"vm1","serial":"None"}', "")
    assert vm.serial == "" and "virtual machine" in vm.note


def test_linux_answer():
    mark = hostinfo.MARK
    output = f"web01\n{mark}\nHPE\n{mark}\nProLiant DL360 Gen10\n{mark}\nMXQ1234ABC\n{mark}\nUbuntu 22.04.4 LTS\n" \
             f"{mark}\n5.15.0-91-generic\n{mark}\naa:bb:cc:00:11:22\n"
    details = hostinfo.parse_linux("10.1.1.20", output)
    assert (details.hostname, details.maker, details.model, details.serial) == ("web01", "HPE",
                                                                                "ProLiant DL360 Gen10", "MXQ1234ABC")
    assert details.os == "Ubuntu 22.04.4 LTS" and details.mac == "AA-BB-CC-00-11-22"
    no_root = hostinfo.parse_linux("10.1.1.21", f"web02\n{mark}\nHPE\n{mark}\nDL360\n{mark}\n\n{mark}\nDebian\n{mark}\n6.1")
    assert no_root.ok and no_root.serial == "" and "root" in no_root.note


def test_phone_web_page():
    page = """<?xml version="1.0" encoding="utf-8"?><DeviceInformation><MACAddress>00AABBCCDDEE</MACAddress>
    <HostName>SEP00AABBCCDDEE</HostName><phoneDN>4001</phoneDN><appLoadID>sip8845_65.14-1-1-0001-136</appLoadID>
    <versionID>sip8845_65.14-1-1-0001-136</versionID><serialNumber>FCH2222ABCD</serialNumber>
    <modelNumber>CP-8845</modelNumber></DeviceInformation>"""
    details = hostinfo.phone_details("10.10.0.22", fetch=lambda url, timeout: page)
    assert details.ok and details.serial == "FCH2222ABCD" and details.model == "CP-8845"
    assert details.mac == "00-AA-BB-CC-DD-EE" and details.software.startswith("sip8845")

    def refused(url, timeout):
        raise OSError("connection refused")
    assert "web access" in hostinfo.phone_details("10.10.0.23", fetch=refused).error
    assert not hostinfo.parse_phone("10.10.0.24", "<html>login</html>").ok


# --------------------------------------------------------------------- The Network Map page's side

def test_read_hardware_again_on_the_map(app, tmp_path, monkeypatch):
    from PyQt5.QtWidgets import QMenu
    from netmap_fakes import FakeNetwork
    from test_assets import crawl as crawl_network, stack_switch
    from test_netmap_tab import Window as MapWindow
    from nomad.netmap import export, store as map_store
    from nomad.netmap.crawl import read_inventory_of
    from nomad.ui import netmap_tab
    monkeypatch.setattr(map_store, "maps_dir", lambda: tmp_path)
    network = FakeNetwork()
    network.add("10.0.0.11", stack_switch())
    network_map = crawl_network(network, read_inventory=False)  # As a map made before NOMAD read hardware
    [key] = network_map.devices
    tab = netmap_tab.NetworkMapTab(MapWindow())
    found = []
    tab.hardware_found.connect(lambda: found.append("found"))
    tab.on_crawled(network_map)
    assert found == ["found"]  # The Assets page brings a crawl in
    tab.read_hardware = lambda settings, address: read_inventory_of(settings, address, client_factory=network.client)
    assert tab.read_hardware_again()
    tab.check_threads[-1].wait(10000)
    QApplication.processEvents()
    device = tab.network_map.devices[key]
    assert device.inventory_read and len(hardware.chassis(device.parts)) == 2
    assert found == ["found", "found"] and "Read the hardware of 1 device" in tab.status_label.text()
    details = netmap_tab.device_html(tab.network_map, key)
    assert "Hardware" in details and "FOC2222B2BB" in details and "17.9.4a" in details
    row = export.device_rows(tab.network_map)[0]
    assert row[export.DEVICE_COLUMNS.index("Serial")] == "FOC1111A1AA, FOC2222B2BB"
    assert row[export.DEVICE_COLUMNS.index("Software")] == "17.9.4a"
    from unittest.mock import Mock
    for name in ("terminal_tab", "scp_tab"):
        saved = Mock()
        saved.saved_matches.return_value = []
        setattr(tab.window, name, saved)
    shown = []
    monkeypatch.setattr(QMenu, "exec_", lambda menu, *args: shown.append(
        {action.text(): action.menu() for action in menu.actions() if not action.isSeparator()}))
    tab.show_device_menu(key, None)
    entries = shown[-1]
    tools = [action.text() for action in entries["Tools"].actions()]
    assert "Read Hardware Again" in tools and "Read Hardware over SSH..." in tools
    assert "Assets" in [action.text() for action in entries["Show In"].actions()]
    tab.shutdown()


def test_details_button_opens_the_dialog_empty(page, monkeypatch):
    tab, store, window = page
    opened = []

    def exec_(dialog):
        opened.append(dialog.addresses_input.toPlainText())
        return False
    monkeypatch.setattr(ReadDetailsDialog, "exec_", exec_)
    tab.details_button.setEnabled(True)
    tab.details_button.click()
    assert opened == [""]


def test_a_mac_finds_its_asset_even_a_port_of_a_stack_member(page, monkeypatch):
    tab, store, window = page
    network_map = window.netmap_tab.network_map
    one, two = hardware.chassis(network_map.devices["sw1"].parts)
    one.update(mac="70-1F-53-00-10-00", mac_source=hardware.BASE, mac_last="70-1F-53-00-10-C7")
    two.update(mac="70-1F-53-00-20-00", mac_source=hardware.BASE, mac_last="70-1F-53-00-20-81")
    bring_in(monkeypatch, tab)
    assert {row[assets_module.COL_MAC] for row in rows(tab.table)} == {"70-1F-53-00-10-00", "70-1F-53-00-20-00",
                                                                       "", "00-AA-BB-CC-DD-EE"}
    tab.search_input.setText("701f.5300.2041")  # Switch 2's port, written as IOS writes it
    assert [row[COL_SERIAL] for row in rows(tab.table)] == ["FOC2"]
    assert "one of its ports'" in tab.status_label.text()
    tab.search_input.setText("00:aa:bb:cc:dd:ee")  # The phone's own
    assert [row[assets_module.COL_MAC] for row in rows(tab.table)] == ["00-AA-BB-CC-DD-EE"]
    tab.search_input.setText("70-1F-53-00-30-00")
    assert rows(tab.table) == [] and "No asset has MAC 70-1F-53-00-30-00" in tab.status_label.text()
    tab.search_input.setText("cafe-bead-acc1")  # A name of hex digits, not written as a MAC is: it's words
    assert tab.table_filter.row_test is None
    tab.search_input.clear()
    assert len(rows(tab.table)) == 4


def test_software_targets_from_the_page(page, monkeypatch):
    from nomad.ui import target_dialogs
    tab, store, window = page
    bring_in(monkeypatch, tab)
    assert not tab.target_list and "target" not in tab.summary_label.text()
    shown = []

    def suggested(dialog):
        shown.append([dialog.table.item(row, 0).text() for row in range(dialog.table.rowCount())])
        return True
    monkeypatch.setattr(target_dialogs.SuggestedTargetsDialog, "exec_", suggested)
    monkeypatch.setattr(target_dialogs.TargetsDialog, "exec_", lambda dialog: dialog.add_suggested() or True)
    tab.show_targets()
    assert shown == [["C9300-48P", "Cisco IP Phone 8845"]]  # Each model without a target, with what it runs
    target = tab.target_list[0]
    assert (target.model, target.target) == ("C9300-48P", "17.9.4a")
    assert "all on target software" in tab.summary_label.text()
    switch = AssetStore(store).by_serial("FOC2")
    AssetStore(store).set_asset(switch.id, software="17.6.5")
    tab.fill_table()
    assert "1 not on target software" in tab.summary_label.text()
    tab.show_combo.setCurrentIndex(tab.show_combo.findData(assets_module.OFF))
    assert [row[COL_SERIAL] for row in rows(tab.table)] == ["FOC2"]
    assert "Behind target: runs 17.6.5" in rows(tab.table)[0][COL_PROBLEMS]
    assert "behind target" in tab.table.item(0, assets_module.COL_SOFTWARE).toolTip()
    tab.select_ids({switch.id})
    assert "Behind target" in tab.details.toPlainText()
    # Right-click > Make 17.6.5 the Target for C9300-48P
    monkeypatch.setattr(target_dialogs.TargetDialog, "exec_", lambda dialog: dialog.save() or True)
    tab.target_from_selected()
    assert tab.target_list[0].target == "17.6.5" and tab.target_list[0].id == target.id
    rows_out = assets_module.export_rows(AssetStore(store).assets(), tab.target_list)
    column = assets_module.EXPORT_COLUMNS.index("On Target")
    assert {row[column] for row in rows_out if row[assets_module.EXPORT_COLUMNS.index("Serial")] == "FOC1"} == {
        "Ahead of target"}


def test_software_targets_overlay_on_the_map(app, tmp_path, monkeypatch):
    from netmap_fakes import FakeNetwork
    from test_assets import crawl as crawl_network, stack_switch
    from test_netmap_tab import Window as MapWindow
    from nomad.ipam.store import SoftwareTarget
    from nomad.netmap import overlays, store as map_store
    from nomad.ui import netmap_tab
    monkeypatch.setattr(map_store, "maps_dir", lambda: tmp_path)
    network = FakeNetwork()
    network.add("10.0.0.11", stack_switch())
    window = MapWindow()
    targets = [SoftwareTarget("t1", "C9300-*", "17.12.1")]

    class Assets:
        def software_targets(self):
            return targets
    window.assets_tab = Assets()
    tab = netmap_tab.NetworkMapTab(window)
    tab.on_crawled(crawl_network(network))
    tab.fill_overlay_menu()
    action = next(action for action in tab.overlay_menu.actions() if action.text() == "Software Targets")
    assert action.isEnabled()
    action.trigger()
    [key] = tab.network_map.devices
    assert tab.overlay_shown == (overlays.SOFTWARE, None)
    assert tab.view.overlay.devices[key].color == "warning" and "1 behind" in tab.vlan_label.text()
    targets[0] = SoftwareTarget("t1", "C9300-*", "17.9.4a")  # Changed on the Assets page
    tab.refresh_software_overlay()
    assert tab.view.overlay.devices[key].color == "success"
    tab.shutdown()


def test_network_devices_read_over_ssh_go_on_the_map_and_the_list(page, monkeypatch):
    from test_inventory_ssh import CAT9K_INVENTORY, CAT9K_VERSION
    from nomad.netmap import inventory_ssh
    tab, store, window = page
    window.netmap_tab.network_map.devices["sw1"].parts = []  # Its SNMP said nothing of its hardware
    applied = []
    window.netmap_tab.apply_ssh_hardware = lambda details: applied.extend(details) or len(details)
    parts, info = inventory_ssh.device_hardware(CAT9K_INVENTORY, CAT9K_VERSION)
    answers = [hostinfo.HostDetails("10.0.0.11", hostinfo.NETWORK, hostname="acc1", software=info["software"],
                                    parts=parts)]
    monkeypatch.setattr(ReadDetailsDialog, "exec_", lambda dialog: True)
    monkeypatch.setattr(ReadDetailsDialog, "answers", lambda dialog: answers)
    monkeypatch.setattr(ReadDetailsDialog, "kind", lambda dialog: "network")
    bring_in(monkeypatch, tab)
    tab.read_details(["10.0.0.11"], hostinfo.NETWORK)
    assert applied == answers  # The map device got it too
    member = AssetStore(store).by_serial("FOC2222B2BB")
    assert member.slot == "Switch 2" and member.map_device == "sw1" and member.software == "17.09.04a"
    assert AssetStore(store).by_serial("LIT2222").parent_id == member.id


def test_read_details_dialog_asks_network_devices_over_ssh(app):
    dialog = ReadDetailsDialog(None, None, ["10.0.0.11"], hostinfo.NETWORK)
    assert dialog.method == hostinfo.NETWORK and dialog.kind() == "network"
    assert dialog.trust_check.isVisibleTo(dialog) and dialog.typed()
    dialog.user_input.setText("admin")
    login = dialog.login()
    assert login["logins"].fallback.username == "admin"
    dialog.deleteLater()
